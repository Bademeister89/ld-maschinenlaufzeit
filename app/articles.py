"""Artikel: je Auftrag und Version (bei Felgen je Nummer) ein Artikel mit kaufmännischen Daten.

Die Artikel entstehen automatisch aus den erfassten Daten (``sync``): Jede Version, die in Läufen oder
Planzeiten eines Auftrags vorkommt, wird ein Artikel – ``21053``, ``21053V1``, ``21053V2``; eine Felge
ist ``10101018``. Umgekehrt legt ein von Hand angelegter Artikel den Auftrag an (``create``).

Je Artikel rechnet ``list_articles``:

- **Laufzeit je Teil:** Ø-Laufzeiten aller Programme der Version über alle Aufspannungen (ohne
  Vorrichtung). Fehlt einem Programm die Ø-Zeit, gilt seine CAM-Planzeit („geschätzt“).
- **Rohling:** Material (Liste mit Dichte und Preis je kg) und Maße – Block L×B×H oder Rund Ø×L in mm.
  Gewicht = Volumen in Litern × Dichte (g/cm³ = kg/l), Materialpreis = Gewicht × Preis je kg.
- **Fräsen:** je Programm Laufzeit je Teil × Stundensatz seiner Maschine (lief es auf mehreren,
  gewichtet nach Laufzeit). Ohne Lauf: Maschine aus der Tebis-Doku, sonst der Durchschnitt aller
  eingetragenen Stundensätze („geschätzt“).
- **Herstellkosten** = Material + Fräsen, **Marge** = VK − Herstellkosten.
- Fertige Teile, Gesamtlaufzeit und letzte Produktion (letzter fertiger Lauf des letzten Programms
  der letzten Aufspannung) kommen aus dem Auftrag.
"""

from __future__ import annotations

import io
import math
import re
import time
from collections import defaultdict
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from . import orders
from .db import Database
from .nc_program import call_name

SHAPES = {"block": "Block", "round": "Rund"}
DIM_MAX_MM = 5000.0
PRICE_MAX = 10_000_000.0
NOTE_MAX = 500
MATERIAL_NAME_MAX = 60
DENSITY_MAX = 25.0  # g/cm³ (Osmium hat 22,6)

_ORDER_KEY = re.compile(r"(\d{4,5})(V\d{1,2})?")
_RIM_KEY = re.compile(r"1\d{7}")


class ArticleError(ValueError):
    """Ungültige Eingabe (Text für die Oberfläche)."""


def article_key(order_key: str, version: str) -> str:
    return f"{order_key}{version}"


def parse_key(text: str) -> tuple[str, str, str]:
    """Artikelnummer → (Auftrag, Version, Art). ``21060``, ``21060v1`` → ("21060", "V1", "order"),
    ``10101018`` → ("10101018", "", "rim")."""
    value = re.sub(r"\s+", "", text or "").upper()
    if _RIM_KEY.fullmatch(value):
        return value, "", "rim"
    match = _ORDER_KEY.fullmatch(value)
    if match:
        return match[1], match[2] or "", "order"
    raise ArticleError(
        f"„{text}“ ist keine Artikelnummer. Beispiele: 21060, 21060V1 (Version) oder 10101018 (Felge)."
    )


# --- Abgleich ------------------------------------------------------------------------------


def sync(db: Database, now: float | None = None) -> int:
    """Für jede Version eines Auftrags, die in Läufen oder Planzeiten vorkommt, einen Artikel anlegen.
    Liefert die Zahl der neuen Artikel. Idempotent."""
    now = time.time() if now is None else now
    kinds = {o["key"]: o["kind"] for o in db.orders()}
    created = 0
    with db.transaction():
        for row in db.order_versions():
            code = orders.parse_program(row["program"])
            if code is None or row["order_key"] not in kinds:
                continue
            version = code.version if kinds[row["order_key"]] == "order" else ""
            created += db.ensure_article(article_key(row["order_key"], version), row["order_key"], version, now)
    return created


def create(db: Database, text: str, now: float | None = None) -> dict[str, Any]:
    """Artikel von Hand anlegen; legt den Auftrag bzw. die Felge mit an, falls es ihn noch nicht gibt."""
    now = time.time() if now is None else now
    order_key, version, kind = parse_key(text)
    key = article_key(order_key, version)
    if db.article(key) is not None:
        raise ArticleError(f"Den Artikel {key} gibt es schon.")
    year = 0 if kind == "rim" else datetime.now().year
    with db.transaction():
        order_created = db.ensure_order(order_key, year, order_key, now, kind=kind) == "created"
        db.ensure_article(key, order_key, version, now, manual=True)
    return {"key": key, "order_key": order_key, "version": version, "kind": kind, "order_created": order_created}


# --- Eingaben prüfen -------------------------------------------------------------------------


def _number(value: Any, label: str, maximum: float, positive: bool = False) -> float | None:
    """Zahl mit Dezimalkomma; leer = None."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    text = str(value).strip().replace("€", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        raise ArticleError(f"{label}: bitte eine Zahl eintragen, z. B. 12,50.") from None
    if not math.isfinite(number) or number < 0 or (positive and number == 0) or number > maximum:
        raise ArticleError(f"{label}: bitte eine Zahl von {'über ' if positive else ''}0 bis {maximum:,.0f} eintragen.".replace(",", "."))
    return number


def validate_article(db: Database, payload: dict[str, Any], current: dict[str, Any]) -> dict[str, Any]:
    """Felder für ``db.update_article``; fehlt ein Schlüssel, bleibt der bisherige Wert."""
    fields = {k: current[k] for k in ("material_id", "shape", "dim_a", "dim_b", "dim_c", "price_ek", "price_vk", "note")}
    if "material_id" in payload:
        value = payload["material_id"]
        if value in (None, ""):
            fields["material_id"] = None
        else:
            try:
                material_id = int(value)
            except (TypeError, ValueError):
                raise ArticleError("Unbekanntes Material.") from None
            if db.material(material_id) is None:
                raise ArticleError("Unbekanntes Material.")
            fields["material_id"] = material_id
    if "shape" in payload:
        shape = payload["shape"] or None
        if shape is not None and shape not in SHAPES:
            raise ArticleError("Form: Block oder Rund.")
        fields["shape"] = shape
    labels = {"dim_a": "Länge bzw. Durchmesser", "dim_b": "Breite bzw. Länge", "dim_c": "Höhe"}
    for name, label in labels.items():
        if name in payload:
            fields[name] = _number(payload[name], f"{label} (mm)", DIM_MAX_MM, positive=True)
    for name, label in (("price_ek", "Preis EK"), ("price_vk", "Preis VK")):
        if name in payload:
            fields[name] = _number(payload[name], label, PRICE_MAX)
    if "note" in payload:
        note = str(payload["note"] or "").strip()
        if len(note) > NOTE_MAX:
            raise ArticleError(f"Die Notiz darf höchstens {NOTE_MAX} Zeichen lang sein.")
        fields["note"] = note
    if fields["shape"] == "round":
        fields["dim_c"] = None  # Rund: Ø und Länge
    return fields


def validate_material(payload: dict[str, Any]) -> tuple[str, float | None, float | None]:
    """(Name, Dichte in g/cm³, Preis je kg)."""
    name = str(payload.get("name") or "").strip()
    if not name:
        raise ArticleError("Bitte einen Namen für das Material eintragen, z. B. Alu 7075.")
    if len(name) > MATERIAL_NAME_MAX:
        raise ArticleError(f"Der Name darf höchstens {MATERIAL_NAME_MAX} Zeichen lang sein.")
    density = _number(payload.get("density"), "Dichte (g/cm³)", DENSITY_MAX, positive=True)
    return name, density, _number(payload.get("price_per_kg"), "Preis je kg", PRICE_MAX)


# --- Rechnen ---------------------------------------------------------------------------------


def volume_l(shape: str | None, a: float | None, b: float | None, c: float | None) -> float | None:
    """Volumen des Rohlings in Litern (Maße in mm)."""
    if shape == "block" and a and b and c:
        return a * b * c / 1e6
    if shape == "round" and a and b:
        return math.pi * (a / 2) ** 2 * b / 1e6
    return None


def _mm(value: float) -> str:
    return f"{value:g}".replace(".", ",")


def blank_text(material: dict[str, Any] | None, shape: str | None, a: float | None, b: float | None, c: float | None) -> str:
    """„Alu 7075 · Block 120 × 80 × 40 mm“ bzw. „Rund Ø 80 × 120 mm“."""
    size = None
    if shape == "block" and a and b and c:
        size = f"Block {_mm(a)} × {_mm(b)} × {_mm(c)} mm"
    elif shape == "round" and a and b:
        size = f"Rund Ø {_mm(a)} × {_mm(b)} mm"
    return " · ".join(p for p in (material["name"] if material else None, size) if p)


def _rate(
    name: str,
    machine_time: dict[str, dict[str, float]],
    plan_machine: str | None,
    machines: list[dict[str, Any]],
) -> tuple[float | None, bool]:
    """Stundensatz für ein Programm: (€/h, geschätzt)."""
    rates = {m["id"]: m["hourly_rate"] for m in machines if m["hourly_rate"] is not None}
    seconds = {mid: s for mid, s in machine_time.get(name, {}).items() if mid in rates and s > 0}
    if seconds:
        return sum(rates[mid] * s for mid, s in seconds.items()) / sum(seconds.values()), False
    if plan_machine:
        hits = [m for m in machines if m["hourly_rate"] is not None and m["name"].lower() in plan_machine.lower()]
        if len(hits) == 1:
            return hits[0]["hourly_rate"], False
    if rates:
        return sum(rates.values()) / len(rates), True
    return None, False


def _version_row(
    version: dict[str, Any] | None,
    machine_time: dict[str, dict[str, float]],
    plans: dict[str, dict[str, Any]],
    machines: list[dict[str, Any]],
    last_finished: dict[str, float],
    base: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Laufzeit je Teil, Fräsen, Stückzahl, Laufzeit und letzte Produktion einer Version. Von der
    Grundversion übernommene Aufspannungen (z. B. die gemeinsame Spannung 1) zählen in Laufzeit je
    Teil und Fräsen mit."""
    setups = [s for s in (version or {}).get("setups", []) if not s["fixture"]]
    inherited = set((version or {}).get("inherited", []))
    shared = [s for s in (base or {}).get("setups", []) if not s["fixture"] and s["setup"] in inherited]
    programs = [p for s in shared + setups for p in s["programs"]]
    part_s: float | None = 0.0 if programs else None
    part_estimated = False
    mill: float | None = 0.0 if programs else None
    mill_estimated = False
    for p in programs:
        seconds = p["avg_run_s"] if p["avg_run_s"] is not None else p["plan_s"]
        if p["avg_run_s"] is None and p["plan_s"] is not None:
            part_estimated = True
        if seconds is None:
            part_s = mill = None
            break
        part_s += seconds
        plan = plans.get(p["call_name"])
        rate, estimated = _rate(p["call_name"], machine_time, plan["machine"] if plan else None, machines)
        if rate is None:
            mill = None
        elif mill is not None:
            mill += seconds / 3600 * rate
            mill_estimated = mill_estimated or estimated
    last = None
    if setups:
        final = max(setups, key=lambda s: s["setup"])["programs"][-1]
        last = last_finished.get(final["call_name"])
    return {
        "inherited": sorted(inherited),
        "part_s": part_s,
        "part_estimated": part_estimated,
        "mill_cost": round(mill, 2) if mill is not None else None,
        "mill_estimated": mill_estimated,
        "parts": (version or {}).get("parts", 0),
        "running_s": (version or {}).get("running_s", 0.0),
        "last_production": last,
    }


def list_articles(db: Database, tz: ZoneInfo, status: str = "all") -> list[dict[str, Any]]:
    """Alle Artikel mit Rohling, Kosten, Preisen und Produktionsdaten (je Auftrag einmal ausgewertet)."""
    sync(db)
    order_rows = {o["key"]: o for o in db.orders(status)}
    materials = {m["id"]: m for m in db.materials()}
    machines = db.machines(include_removed=True)
    by_order: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for art in db.articles():
        if art["order_key"] in order_rows:
            by_order[art["order_key"]].append(art)
    rows = []
    for order_key, arts in by_order.items():
        detail = orders.order_detail(db, order_key, tz)
        if detail is None:
            continue
        machine_time: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        for r in db.machine_run_time(order_key):
            machine_time[call_name(r["program"])][r["machine_id"]] += r["seconds"] or 0.0
        last_finished: dict[str, float] = {}
        for r in db.last_finished(order_key):
            name = call_name(r["program"])
            last_finished[name] = max(last_finished.get(name, 0.0), r["last"] or 0.0)
        plans = {p["name"]: p for p in db.plans_of_order(order_key)}
        versions = {v["version"]: v for v in detail["versions"]}
        order = detail["order"]
        for art in arts:
            if versions.get(art["version"], {}).get("pre_stage"):
                continue  # Grundversion ist nur die gemeinsame Vorstufe (Spannung 1) der Versionen
            material = materials.get(art["material_id"])
            volume = volume_l(art["shape"], art["dim_a"], art["dim_b"], art["dim_c"])
            density = material["density"] if material else None
            weight = volume * density if volume is not None and density is not None else None  # g/cm³ = kg/l
            price_kg = material["price_per_kg"] if material else None
            material_cost = round(weight * price_kg, 2) if weight is not None and price_kg is not None else None
            prod = _version_row(versions.get(art["version"]), machine_time, plans, machines, last_finished, versions.get(""))
            known = [c for c in (material_cost, prod["mill_cost"]) if c is not None]
            cost = round(sum(known), 2) if known else None
            cost_complete = material_cost is not None and prod["mill_cost"] is not None
            vk = art["price_vk"]
            margin = round(vk - cost, 2) if vk is not None and cost is not None and cost_complete else None
            rows.append({
                "key": art["key"],
                "order_key": order_key,
                "version": art["version"],
                "kind": order["kind"],
                "title": order["title"],
                "rim": order.get("rim"),
                "status": order["status"],
                "image": order["image"],
                "manual": bool(art["manual"]),
                "material_id": art["material_id"],
                "material": material["name"] if material else None,
                "shape": art["shape"],
                "dim_a": art["dim_a"],
                "dim_b": art["dim_b"],
                "dim_c": art["dim_c"],
                "blank": blank_text(material, art["shape"], art["dim_a"], art["dim_b"], art["dim_c"]),
                "volume_l": volume,
                "density": density,
                "weight_kg": weight,
                "price_per_kg": price_kg,
                "material_cost": material_cost,
                **prod,
                "cost": cost,
                "cost_complete": cost_complete,
                "price_ek": art["price_ek"],
                "price_vk": vk,
                "margin": margin,
                "margin_pct": round(margin / vk * 100, 1) if margin is not None and vk else None,
                "note": art["note"],
            })
    return sorted(rows, key=lambda r: (r["order_key"], orders.version_order(r["version"])))


# --- Excel-Export ------------------------------------------------------------------------------
# Eine Arbeitsmappe zum Weiterrechnen: Rohling, Material, Stundensatz und Preise sind Eingaben (gelb),
# Volumen, Gewicht, Materialpreis, Fräsen, Herstellkosten und Marge sind Formeln wie in der Artikelliste.
# Ändert man in Excel z. B. die Maße oder das Material, rechnet die Zeile neu (Stand des Exports).

EUR = '#,##0.00 "€"'
DURATION = "[h]:mm"
FIRST_ROW = 3  # Zeile 1: Gruppen, Zeile 2: Spaltenköpfe
THUMB_PX = (96, 72)  # Vorschaubild in der Tabelle (4:3)
SHAPE_NAMES = {"block": "Block", "round": "Rund"}

# (Schlüssel, Überschrift, Breite, Format, Eingabe?)
_COLS = [
    ("image", "Bild", 14, None, False),
    ("key", "Artikel", 12, "@", False),
    ("title", "Bezeichnung", 30, "@", False),
    ("kind", "Art", 9, "@", False),
    ("status", "Status", 13, "@", False),
    ("material", "Material", 16, "@", True),
    ("shape", "Form", 9, "@", True),
    ("dim_a", "Maß 1 (mm)\nL bzw. Ø", 11, "0.0", True),
    ("dim_b", "Maß 2 (mm)\nB bzw. L", 11, "0.0", True),
    ("dim_c", "Maß 3 (mm)\nH (Block)", 11, "0.0", True),
    ("volume", "Volumen (l)", 11, "0.000", False),
    ("density", "Dichte (g/cm³)", 11, "0.00", False),
    ("weight", "Gewicht (kg)", 11, "0.000", False),
    ("price_kg", "€ je kg", 11, EUR, False),
    ("material_cost", "Material­preis", 13, EUR, False),
    ("part", "Laufzeit je Teil", 12, DURATION, True),
    ("rate", "Stundensatz (€/h)", 12, EUR, True),
    ("mill", "Preis Fräsen", 13, EUR, False),
    ("cost", "Herstell­kosten", 14, EUR, False),
    ("ek", "Preis EK", 12, EUR, True),
    ("vk", "Preis VK", 12, EUR, True),
    ("margin", "Marge", 13, EUR, False),
    ("margin_pct", "Marge %", 10, "0.0%", False),
    ("parts", "Stück produziert", 11, "0", False),
    ("running", "Gesamt­laufzeit", 12, DURATION, False),
    ("last", "Letzte Produktion", 13, "DD.MM.YYYY", False),
    ("hints", "Hinweise", 34, "@", False),
    ("note", "Notiz", 30, "@", False),
]
_GROUPS = [
    ("Artikel", "image", "status", "2F4F6F"),
    ("Rohling und Material", "material", "material_cost", "3E6B48"),
    ("Fertigung", "part", "mill", "6B4E2F"),
    ("Kosten und Preise", "cost", "margin_pct", "5B3E6B"),
    ("Produktion", "parts", "last", "2F5F6B"),
    ("", "hints", "note", "4A4A48"),
]


def hints(row: dict[str, Any]) -> str:
    notes = []
    if row.get("inherited"):
        notes.append(f"inkl. Spannung {', '.join(map(str, row['inherited']))} der Grundversion")
    if row["part_estimated"]:
        notes.append("Laufzeit teils aus CAM-Planzeit")
    if row["mill_estimated"]:
        notes.append("Stundensatz geschätzt (Durchschnitt)")
    if row["part_s"] is not None and row["mill_cost"] is None:
        notes.append("Stundensatz fehlt")
    return "; ".join(notes)


def _jpeg_size(data: bytes) -> tuple[int, int] | None:
    """Breite und Höhe aus dem JPEG-Kopf (SOF-Segment) – ohne Bildbibliothek."""
    i = 2
    while i + 9 < len(data) and data[i] == 0xFF:
        marker, length = data[i + 1], int.from_bytes(data[i + 2 : i + 4], "big")
        if 0xC0 <= marker <= 0xCF and marker not in (0xC4, 0xC8, 0xCC):
            return int.from_bytes(data[i + 7 : i + 9], "big"), int.from_bytes(data[i + 5 : i + 7], "big")
        i += 2 + length
    return None


def _image(data: bytes) -> Any:
    """Bild für openpyxl ohne Pillow: die Vorschaubilder sind schon JPEG und werden unverändert eingebettet."""
    from openpyxl.drawing.image import Image

    class JpegImage(Image):
        def __init__(self, raw: bytes):  # noqa: D107 – Image.__init__ bräuchte Pillow
            self.ref, self._raw, self.format = raw, raw, "jpeg"
            self.width, self.height = THUMB_PX

        def _data(self) -> bytes:
            return self._raw

    return JpegImage(data)


def export_xlsx(
    rows: list[dict[str, Any]],
    materials: list[dict[str, Any]],
    tz: ZoneInfo,
    now: float,
    images: dict[str, bytes] | None = None,
) -> bytes:
    """Artikelliste als Excel-Arbeitsmappe mit Bildern und Formeln (Blatt „Artikel“), der Materialliste
    (Blatt „Materialien“, Grundlage der Formeln) und einer kurzen Erläuterung."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    images = images or {}
    col = {key: get_column_letter(i) for i, (key, *_) in enumerate(_COLS, start=1)}
    thin = Side(style="thin", color="C9C7BE")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    input_fill = PatternFill("solid", fgColor="FFF6D6")  # Eingaben: hellgelb
    zebra = PatternFill("solid", fgColor="F4F4F1")
    white = Font(bold=True, color="FFFFFF", size=12)

    wb = Workbook()
    ws = wb.active
    ws.title = "Artikel"
    ws.sheet_properties.tabColor = "2F4F6F"

    # Zeile 1: Gruppen über mehrere Spalten, Zeile 2: Spaltenköpfe in der Farbe der Gruppe
    group_of = {}
    for title, first, last, color in _GROUPS:
        a, b = col[first], col[last]
        ws.merge_cells(f"{a}1:{b}1")
        cell = ws[f"{a}1"]
        cell.value = title or None
        cell.font, cell.fill = white, PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        start, end = [k for k, *_ in _COLS].index(first), [k for k, *_ in _COLS].index(last)
        for key, *_ in _COLS[start : end + 1]:
            group_of[key] = color
    for i, (key, title, width, _, is_input) in enumerate(_COLS, start=1):
        cell = ws.cell(row=2, column=i, value=title.replace("­", ""))
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=group_of[key])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = grid
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 42

    materials_ref = f"Materialien!$A$2:$C${max(len(materials) + 1, 2) + 200}"
    for n, row in enumerate(rows):
        r = FIRST_ROW + n
        c = {key: f"{col[key]}{r}" for key in col}
        rate = None
        if row["mill_cost"] is not None and row["part_s"]:
            rate = round(row["mill_cost"] / (row["part_s"] / 3600), 2)  # wirksamer Stundensatz (gewichtet)
        last = datetime.fromtimestamp(row["last_production"], tz).replace(tzinfo=None) if row["last_production"] else None
        values = {
            "key": row["key"],
            "title": row["title"] or (row["rim"] or {}).get("label") or "",
            "kind": "Felge" if row["kind"] == "rim" else "Auftrag",
            "status": "abgeschlossen" if row["status"] == "closed" else "offen",
            "material": row["material"],
            "shape": SHAPE_NAMES.get(row["shape"] or ""),
            "dim_a": row["dim_a"],
            "dim_b": row["dim_b"],
            "dim_c": row["dim_c"],
            # Formeln wie in der Artikelliste
            "volume": (
                f'=IF(AND({c["shape"]}="Block",N({c["dim_a"]})>0,N({c["dim_b"]})>0,N({c["dim_c"]})>0),'
                f'{c["dim_a"]}*{c["dim_b"]}*{c["dim_c"]}/1000000,'
                f'IF(AND({c["shape"]}="Rund",N({c["dim_a"]})>0,N({c["dim_b"]})>0),'
                f'PI()*({c["dim_a"]}/2)^2*{c["dim_b"]}/1000000,""))'
            ),
            "density": f'=IFERROR(VLOOKUP({c["material"]},{materials_ref},2,FALSE),"")',
            "weight": f'=IF(AND(ISNUMBER({c["volume"]}),ISNUMBER({c["density"]})),{c["volume"]}*{c["density"]},"")',
            "price_kg": f'=IFERROR(VLOOKUP({c["material"]},{materials_ref},3,FALSE),"")',
            "material_cost": (
                f'=IF(AND(ISNUMBER({c["weight"]}),ISNUMBER({c["price_kg"]})),ROUND({c["weight"]}*{c["price_kg"]},2),"")'
            ),
            "part": row["part_s"] / 86400 if row["part_s"] is not None else None,
            "rate": rate,
            "mill": f'=IF(AND(ISNUMBER({c["part"]}),ISNUMBER({c["rate"]})),ROUND({c["part"]}*24*{c["rate"]},2),"")',
            "cost": f'=IF(COUNT({c["material_cost"]},{c["mill"]})=0,"",SUM({c["material_cost"]},{c["mill"]}))',
            "ek": row["price_ek"],
            "vk": row["price_vk"],
            "margin": (
                f'=IF(AND(ISNUMBER({c["vk"]}),COUNT({c["material_cost"]},{c["mill"]})=2),{c["vk"]}-{c["cost"]},"")'
            ),
            "margin_pct": f'=IF(AND(ISNUMBER({c["margin"]}),N({c["vk"]})>0),{c["margin"]}/{c["vk"]},"")',
            "parts": row["parts"],
            "running": row["running_s"] / 86400,
            "last": last,
            "hints": hints(row) or None,
            "note": row["note"] or None,
        }
        for i, (key, _, _, fmt, is_input) in enumerate(_COLS, start=1):
            cell = ws.cell(row=r, column=i, value=values.get(key))
            if fmt:
                cell.number_format = fmt
            cell.border = grid
            cell.font = Font(size=11, bold=key in ("key", "cost", "margin"))
            wrap = key in ("title", "hints", "note")
            cell.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="left" if fmt == "@" else None)
            if is_input:
                cell.fill = input_fill
            elif n % 2:
                cell.fill = zebra
        ws.row_dimensions[r].height = 58
        data = images.get(row["key"])
        if data:
            img = _image(data)
            img.anchor = c["image"]
            ws.add_image(img)

    last_row = FIRST_ROW + len(rows) - 1
    if rows:
        total = last_row + 1
        ws.cell(row=total, column=2, value="Summe").font = Font(bold=True, size=11)
        for key in ("parts", "running"):
            cell = ws[f"{col[key]}{total}"]
            cell.value = f"=SUM({col[key]}{FIRST_ROW}:{col[key]}{last_row})"
            cell.font = Font(bold=True, size=11)
            cell.number_format = dict((k, f) for k, _, _, f, _ in _COLS)[key]
        for i in range(1, len(_COLS) + 1):
            ws.cell(row=total, column=i).border = Border(top=Side(style="medium", color="2F4F6F"))
        ws.row_dimensions[total].height = 24
        # Auswahllisten für die Eingaben
        shapes = DataValidation(type="list", formula1='"Block,Rund"', allow_blank=True)
        shapes.add(f"{col['shape']}{FIRST_ROW}:{col['shape']}{last_row}")
        names = DataValidation(type="list", formula1=f"=Materialien!$A$2:$A${len(materials) + 201}", allow_blank=True)
        names.error, names.errorTitle = "Bitte ein Material aus dem Blatt „Materialien“ wählen.", "Material"
        names.add(f"{col['material']}{FIRST_ROW}:{col['material']}{last_row}")
        ws.add_data_validation(shapes)
        ws.add_data_validation(names)
    ws.freeze_panes = f"{col['title']}{FIRST_ROW}"  # Bild und Artikel bleiben beim Scrollen stehen
    ws.auto_filter.ref = f"A2:{get_column_letter(len(_COLS))}{max(last_row, 2)}"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"

    ms = wb.create_sheet("Materialien")
    ms.sheet_properties.tabColor = "3E6B48"
    ms.append(["Material", "Dichte (g/cm³)", "Preis je kg", "Artikel"])
    for i, width in enumerate((26, 16, 16, 10), start=1):
        cell = ms.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="3E6B48")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ms.column_dimensions[get_column_letter(i)].width = width
    ms.row_dimensions[1].height = 24
    for m in materials:
        ms.append([m["name"], m["density"], m["price_per_kg"], m["articles"]])
        r = ms.max_row
        ms.cell(row=r, column=2).number_format = "0.00"
        ms.cell(row=r, column=3).number_format = EUR
        for i in (1, 2, 3):
            ms.cell(row=r, column=i).fill = input_fill
            ms.cell(row=r, column=i).border = grid
    ms.freeze_panes = "A2"

    info = wb.create_sheet("Erläuterung")
    lines = [
        ("Artikelliste – Stand des Exports", True),
        (f"Erstellt am {datetime.fromtimestamp(now, tz):%d.%m.%Y um %H:%M} Uhr mit LD-Machine-Viewer.", False),
        ("", False),
        ("Gelbe Zellen sind Eingaben, alle anderen rechnen mit Formeln wie in der Artikelliste:", True),
        ("Volumen (l) = Block: L × B × H / 1.000.000 · Rund: π × (Ø/2)² × L / 1.000.000 (Maße in mm)", False),
        ("Dichte und Preis je kg kommen über das Material aus dem Blatt „Materialien“ (dort änderbar).", False),
        ("Gewicht (kg) = Volumen × Dichte · Materialpreis = Gewicht × Preis je kg", False),
        ("Preis Fräsen = Laufzeit je Teil × Stundensatz (bei mehreren Maschinen nach Laufzeit gewichtet)", False),
        ("Herstellkosten = Materialpreis + Preis Fräsen · Marge = Preis VK − Herstellkosten", False),
        ("", False),
        ("Laufzeit je Teil: Summe der Ø-Laufzeiten aller Aufspannungen, bei Versionen inkl. der gemeinsamen", False),
        ("Spannung 1 der Grundversion. Ohne fertigen Lauf gilt die CAM-Planzeit (siehe Hinweise).", False),
        ("Änderungen hier wirken nur in dieser Datei, nicht in der App.", False),
    ]
    for text, bold in lines:
        info.append([text])
        info.cell(row=info.max_row, column=1).font = Font(bold=bold, size=12 if bold else 11)
    info.column_dimensions["A"].width = 110

    wb.calculation.fullCalcOnLoad = True  # Excel rechnet die Formeln beim Öffnen
    wb.properties.creator = "LD-Machine-Viewer"
    wb.properties.created = datetime.fromtimestamp(now, tz).replace(tzinfo=None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
