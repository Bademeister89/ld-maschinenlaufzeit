"""Artikelverwaltung: Abgleich mit den Aufträgen, Rohling und Kosten, Preise, Excel-Export."""

import io
import math
import sqlite3
from datetime import timedelta
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import articles
from app.config import MachineConfig, Settings
from app.db import SCHEMA_VERSION, Database
from app.main import create_app

from .conftest import feed, snap

TZ = ZoneInfo("Europe/Berlin")
P11 = "TNC:\\X\\26-21055-01-01.H"
P21 = "TNC:\\X\\26-21055-02-01.H"
V1 = "TNC:\\X\\26-21055V1-01-01.H"


def run_part(col, t0, program, run_s):
    feed(col, (t0, snap("IDLE", program)), (t0 + 10, snap("STARTED", program)), (t0 + 10 + run_s, snap("FINISHED", program)))
    return t0 + 10 + run_s


@pytest.fixture
def client(tmp_path):
    settings = Settings(
        machines=(MachineConfig("m1", "DMU 70", "10.0.0.1"), MachineConfig("m2", "DMU 105", "10.0.0.2")),
        db_path=tmp_path / "a.db",
        simulate=True,
    )
    app = create_app(settings, run_collectors=False)
    with TestClient(app) as c:
        yield c


def rate(client, machine_id, value):
    m = next(x for x in client.get("/api/config").json()["machines"] if x["id"] == machine_id)
    r = client.put(f"/api/config/machines/{machine_id}", json={**m, "hourly_rate": value})
    assert r.status_code == 200, r.text
    return r.json()


def article(client, key):
    return next(a for a in client.get("/api/articles").json()["articles"] if a["key"] == key)


# --- Nummern, Volumen ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [("21060", ("21060", "", "order")), (" 21060v1 ", ("21060", "V1", "order")), ("4711V12", ("4711", "V12", "order")),
     ("10101018", ("10101018", "", "rim"))],
)
def test_parse_key(text, expected):
    assert articles.parse_key(text) == expected


@pytest.mark.parametrize("text", ["", "123", "21060X", "21060V", "2106001-01", "20101018"])
def test_parse_key_rejects(text):
    with pytest.raises(articles.ArticleError):
        articles.parse_key(text)


def test_volume_and_blank_text():
    assert articles.volume_l("block", 120, 80, 40) == pytest.approx(0.384)
    assert articles.volume_l("round", 80, 120, None) == pytest.approx(math.pi * 40**2 * 120 / 1e6)
    assert articles.volume_l("block", 120, 80, None) is None
    assert articles.blank_text({"name": "Alu 7075"}, "block", 120, 80, 40.5) == "Alu 7075 · Block 120 × 80 × 40,5 mm"
    assert articles.blank_text(None, "round", 80, 120, None) == "Rund Ø 80 × 120 mm"


# --- Abgleich -------------------------------------------------------------------------------------


def test_sync_creates_one_article_per_version_and_manual_articles_create_the_order(client):
    ctx = client.app.state.ctx
    col = ctx.collectors["m1"]
    t = run_part(col, 1000, P11, 600)
    t = run_part(col, t + 10, V1, 300)
    client.put("/api/orders/21055/plans", json={"program": "26-21055V2-01-01", "time": "1"})  # V2 nur geplant
    keys = [a["key"] for a in client.get("/api/articles").json()["articles"]]
    assert keys == ["21055", "21055V1", "21055V2"]
    assert articles.sync(ctx.db) == 0  # idempotent

    r = client.post("/api/articles", json={"key": "21060v1"})
    assert r.status_code == 201 and r.json()["order_created"] is True
    assert ctx.db.order("21060") is not None
    a = article(client, "21060V1")
    assert (a["manual"], a["parts"], a["part_s"]) == (True, 0, None)
    assert client.post("/api/articles", json={"key": "21060V1"}).status_code == 400  # gibt es schon
    assert client.post("/api/articles", json={"key": "abc"}).status_code == 400
    assert client.post("/api/articles", json={"key": "10101018"}).json()["kind"] == "rim"

    # Auftrag löschen: Artikel weg; Stammdaten entfernen: Abgleich legt ihn neu an (ohne Preise)
    client.put("/api/articles/21055V1", json={"price_vk": "99"})
    assert client.delete("/api/articles/21055V1").status_code == 204
    assert article(client, "21055V1")["price_vk"] is None
    assert client.delete("/api/orders/21060").status_code == 200
    assert "21060V1" not in [a["key"] for a in client.get("/api/articles").json()["articles"]]


# --- Kosten --------------------------------------------------------------------------------------


def test_costs_material_milling_and_margin(client):
    ctx = client.app.state.ctx
    t = run_part(ctx.collectors["m1"], 1000, P11, 3600)  # Spannung 1: 1 h an der DMU 70
    run_part(ctx.collectors["m2"], t + 10, P21, 1800)  # Spannung 2: 0,5 h an der DMU 105
    rate(client, "m1", "90")
    rate(client, "m2", "120,50")
    mat = client.post("/api/config/materials", json={"name": "Alu 7075", "density": "2,81", "price_per_kg": "5,20"}).json()
    r = client.put("/api/articles/21055", json={
        "material_id": mat["id"], "shape": "block", "dim_a": "120", "dim_b": "80", "dim_c": "40",
        "price_ek": "100", "price_vk": "250,00",
    })
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["blank"] == "Alu 7075 · Block 120 × 80 × 40 mm"
    # 0,384 l × 2,81 kg/l = 1,079 kg × 5,20 €/kg
    assert (a["volume_l"], a["density"], a["price_per_kg"]) == (pytest.approx(0.384), 2.81, 5.2)
    assert (a["weight_kg"], a["material_cost"]) == (pytest.approx(1.07904), 5.61)
    assert a["part_s"] == pytest.approx(5400)
    assert a["mill_cost"] == pytest.approx(90 + 60.25)
    assert (a["cost"], a["cost_complete"]) == (pytest.approx(155.86), True)
    assert (a["margin"], a["margin_pct"]) == (pytest.approx(94.14), 37.7)
    assert (a["parts"], a["mill_estimated"], a["part_estimated"]) == (1, False, False)
    assert a["last_production"] is not None

    # Preis je kg geändert → Materialpreis sofort neu; ohne Dichte kein Gewicht und kein Preis
    client.put(f"/api/config/materials/{mat['id']}", json={"name": "Alu 7075", "density": "2,81", "price_per_kg": "10"})
    assert article(client, "21055")["material_cost"] == pytest.approx(10.79)
    client.put(f"/api/config/materials/{mat['id']}", json={"name": "Alu 7075", "price_per_kg": "10"})
    assert (article(client, "21055")["weight_kg"], article(client, "21055")["material_cost"]) == (None, None)
    # Material entfernt: Maße bleiben, Herstellkosten nur noch „mind.“
    assert client.delete(f"/api/config/materials/{mat['id']}").status_code == 204
    a = article(client, "21055")
    assert (a["material"], a["shape"], a["cost_complete"], a["margin"]) == (None, "block", False, None)


def test_milling_rate_weighted_and_plan_fallback(client):
    ctx = client.app.state.ctx
    t = run_part(ctx.collectors["m1"], 1000, P11, 3000)
    run_part(ctx.collectors["m2"], t + 10, P11, 1000)  # dasselbe Programm an der anderen Maschine
    client.put("/api/orders/21055/plans", json={"program": "26-21055-02-01", "time": "1"})  # nie gelaufen
    a = article(client, "21055")
    assert a["mill_cost"] is None and a["part_s"] == pytest.approx(2000 + 3600) and a["part_estimated"]
    rate(client, "m1", "100")
    rate(client, "m2", "60")
    a = article(client, "21055")
    # Ø 2000 s × (100·3000 + 60·1000)/4000 = 90 €/h → 50 €; Planzeit 1 h × Durchschnitt 80 €/h (geschätzt)
    assert a["mill_cost"] == pytest.approx(50 + 80)
    assert a["mill_estimated"] is True


def test_article_validation(client):
    run_part(client.app.state.ctx.collectors["m1"], 1000, P11, 60)
    for payload in ({"dim_a": "-1"}, {"dim_a": "0"}, {"dim_a": "9999"}, {"price_vk": "abc"}, {"price_ek": "-5"},
                    {"shape": "kugel"}, {"material_id": 99}, {"note": "x" * 501}):
        assert client.put("/api/articles/21055", json=payload).status_code == 400, payload
    r = client.put("/api/articles/21055", json={"shape": "round", "dim_a": "80,5", "dim_b": "120", "dim_c": "5", "price_vk": "1.234,50"})
    assert (r.json()["dim_a"], r.json()["dim_c"], r.json()["price_vk"]) == (80.5, None, 1234.5)
    assert client.put("/api/articles/99999", json={}).status_code == 404
    assert client.post("/api/config/materials", json={"name": ""}).status_code == 400
    for bad in ({"name": "Stahl", "density": "0"}, {"name": "Stahl", "density": "30"}, {"name": "Stahl", "price_per_kg": "x"}):
        assert client.post("/api/config/materials", json=bad).status_code == 400
    client.post("/api/config/materials", json={"name": "Stahl", "density": "7,85"})
    assert client.post("/api/config/materials", json={"name": "stahl"}).status_code == 409
    m = next(x for x in client.get("/api/config").json()["machines"] if x["id"] == "m1")
    assert client.put("/api/config/machines/m1", json={**m, "hourly_rate": "viel"}).status_code == 400


# --- Excel-Export ---------------------------------------------------------------------------------


def test_excel_export(client):
    ctx = client.app.state.ctx
    t = run_part(ctx.collectors["m1"], 1000, P11, 3600)
    run_part(ctx.collectors["m1"], t + 10, V1, 1800)
    rate(client, "m1", "90")
    mat = client.post("/api/config/materials", json={"name": "Alu 7075", "density": "2,81", "price_per_kg": "5,20"}).json()
    client.put("/api/articles/21055", json={"material_id": mat["id"], "shape": "block", "dim_a": "100", "dim_b": "100", "dim_c": "100", "price_vk": "200"})
    r = client.get("/api/articles/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["content-disposition"].startswith('attachment; filename="artikel_')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Artikel", "Materialien"]
    ws = wb["Artikel"]
    head = [c.value for c in ws[1]]
    assert head[:4] == ["Artikel", "Art", "Bezeichnung", "Status"]
    col = {name: i for i, name in enumerate(head)}
    row = [c for c in ws[2]]
    assert row[col["Artikel"]].value == "21055"
    assert row[col["Gewicht (kg)"]].value == pytest.approx(2.81)  # 1 l Alu 7075
    assert row[col["Materialpreis"]].value == pytest.approx(14.61)
    assert row[col["Materialpreis"]].number_format == '#,##0.00 "€"'
    assert row[col["Laufzeit je Teil"]].value == timedelta(hours=1)  # Excel-Dauer, als Zeit eingelesen
    assert row[col["Laufzeit je Teil"]].number_format == "[h]:mm"
    assert row[col["Preis Fräsen"]].value == pytest.approx(90)
    assert row[col["Marge"]].value == pytest.approx(200 - 104.61)
    assert row[col["Letzte Produktion"]].number_format == "DD.MM.YYYY"
    assert ws.cell(row=3, column=1).value == "21055V1"
    assert ws.cell(row=4, column=1).value == "Summe"
    assert ws.cell(row=4, column=col["Stück produziert"] + 1).value.startswith("=SUM(")
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref.startswith("A1:")
    assert [c.value for c in wb["Materialien"][2]] == ["Alu 7075", 2.81, 5.2, 1]


# --- Migration ----------------------------------------------------------------------------------


def test_update_from_schema_16_adds_articles_and_hourly_rate(tmp_path):
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.executescript(
        "CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);"
        "INSERT INTO meta VALUES ('schema_version', '16');"
        "CREATE TABLE machines (id TEXT PRIMARY KEY, name TEXT NOT NULL, host TEXT NOT NULL, port INTEGER NOT NULL);"
        "INSERT INTO machines VALUES ('m1', 'DMU 70', '10.0.0.1', 19000);"
    )
    con.close()
    db = Database(path)
    assert db.get_meta("schema_version") == str(SCHEMA_VERSION) == "17"
    assert db.machine("m1")["hourly_rate"] is None
    assert db.articles() == [] and db.materials() == []
    db.close()
